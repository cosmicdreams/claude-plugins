"""Build the four-slide Design Lab science fair presentation."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "design-lab-science-fair.pptx"
FIGMA_LOGO = ROOT / "assets" / "figma-logo.png"
VELIR_LOGO = ROOT.parent / "drover" / "assets" / "branding" / "velir-logo.png"

NAVY = RGBColor(0x00, 0x1B, 0x67)
BLUE = RGBColor(0x00, 0x51, 0xFF)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x55, 0x73, 0x82)
PALE = RGBColor(0xE6, 0xE8, 0xFF)


def text(slide, value, x, y, w, h, size, color, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0)
    frame.margin_top = frame.margin_bottom = Inches(0)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = "IBM Plex Sans"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return shape


def add_velir_logo(slide):
    slide.shapes.add_picture(str(VELIR_LOGO), Inches(11.48), Inches(6.82), width=Inches(1.18))


prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

# 1 — opener
s = prs.slides.add_slide(blank)
s.background.fill.solid()
s.background.fill.fore_color.rgb = NAVY
text(s, "Design Lab", 0.72, 2.08, 11.5, 1.05, 64, WHITE, True)
text(s, "Code → Figma component library generator", 0.74, 3.38, 11.5, 0.75, 29, WHITE)
text(s, "VELIR  /  INTERNAL SCIENCE FAIR", 0.76, 6.78, 7.0, 0.30, 12, WHITE, True)
s.notes_slide.notes_text_frame.text = "Introduce Design Lab as the code-to-Figma component library generator."

# 2 — video-ready pre-show
s = prs.slides.add_slide(blank)
s.background.fill.solid()
s.background.fill.fore_color.rgb = WHITE
text(s, "Pre-show", 0.72, 0.35, 11.8, 0.62, 37, BLUE, True)
frame = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.64), Inches(1.19), Inches(10.05), Inches(5.65))
frame.fill.solid()
frame.fill.fore_color.rgb = NAVY
frame.line.color.rgb = PALE
frame.line.width = Pt(1)
frame.adjustments[0] = 0.03
frame._element.find(qn("p:style")).find(qn("a:effectRef")).set("idx", "0")
text(s, "LOOM VIDEO", 2.55, 3.15, 8.2, 0.55, 28, WHITE, True, PP_ALIGN.CENTER)
text(s, "Generating the Figma file from scratch", 2.15, 3.77, 9.0, 0.50, 22, WHITE, False, PP_ALIGN.CENTER)
add_velir_logo(s)
s.notes_slide.notes_text_frame.text = (
    "Replace the navy 16:9 area with the Loom recording before the event. "
    "In PowerPoint, use Insert > Video > Video on My PC (or Online Video), "
    "then size the video to cover the navy area. For unattended pre-show, "
    "set playback to Start Automatically and Loop until Stopped."
)

# 3 — live Figma walkthrough
s = prs.slides.add_slide(blank)
s.background.fill.solid()
s.background.fill.fore_color.rgb = WHITE
text(s, "Figma walkthrough", 0.72, 0.35, 11.8, 0.62, 37, BLUE, True)
logo = s.shapes.add_picture(str(FIGMA_LOGO), Inches(5.55), Inches(1.45), height=Inches(3.85))
logo.click_action.hyperlink.address = "https://www.figma.com/files/"
launch = text(s, "Open Figma", 4.65, 5.64, 4.0, 0.57, 27, BLUE, True, PP_ALIGN.CENTER)
launch.click_action.hyperlink.address = "https://www.figma.com/files/"
add_velir_logo(s)
s.notes_slide.notes_text_frame.text = (
    "Click the Figma logo or Open Figma to open Figma in a browser. "
    "Replace both hyperlinks with the direct component-library file URL when it is available. "
    "Figma logo source: https://commons.wikimedia.org/wiki/File:Figma-logo.svg"
)

# 4 — questions
s = prs.slides.add_slide(blank)
s.background.fill.solid()
s.background.fill.fore_color.rgb = NAVY
text(s, "Questions?", 0.72, 2.20, 11.6, 1.08, 64, WHITE, True)
text(s, "Get in the comments!", 0.75, 3.52, 11.5, 0.72, 30, WHITE)
text(s, "VELIR  /  DESIGN LAB", 0.76, 6.78, 7.0, 0.30, 12, WHITE, True)
s.notes_slide.notes_text_frame.text = "Invite questions and ask attendees to add feedback in the event comments."

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUTPUT)
print(OUTPUT)
